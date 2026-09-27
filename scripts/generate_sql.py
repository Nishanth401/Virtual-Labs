import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\erohi\OneDrive\Desktop\Virtual-Labs\AIDS student count.xlsx')

cohort_info = {
    'II AIDS': {'year': 'II Year', 'cohort': 'II AIDS', 'year_sem': 'Year II / Semester IV', 'class': 'II AIDS'},
    'III AIDS': {'year': 'III Year', 'cohort': 'III AIDS', 'year_sem': 'Year III / Semester VI', 'class': 'III AIDS'},
    'IV AIDS': {'year': 'IV Year', 'cohort': 'IV AIDS', 'year_sem': 'Year IV / Semester VIII', 'class': 'IV AIDS'}
}

def escape_sql(val):
    if val is None:
        return 'NULL'
    s = str(val).replace("'", "''")
    return f"'{s}'"

student_rows = []
profile_rows = []
auth_user_rows = []
auth_identity_rows = []

for sheetname, meta in cohort_info.items():
    sheet = wb[sheetname]
    for r in range(2, sheet.max_row + 1):
        sid = sheet.cell(r, 1).value
        sno = sheet.cell(r, 2).value
        essl = sheet.cell(r, 3).value
        name = sheet.cell(r, 4).value
        reg = sheet.cell(r, 6).value
        advisor = sheet.cell(r, 7).value
        if name and reg:
            reg_clean = str(reg).strip().replace('.0', '')
            name_clean = str(name).strip().upper()
            adv_clean = str(advisor).strip() if advisor is not None else 'Faculty Advisor'
            essl_clean = str(essl).strip() if essl is not None else ''
            sid_val = int(sid) if sid is not None and str(sid).isdigit() else 'NULL'
            sno_val = int(sno) if sno is not None and str(sno).isdigit() else 'NULL'
            email = f'{reg_clean}@vsb.ac.in'

            student_rows.append(
                f"({sid_val}, {sno_val}, {escape_sql(essl_clean)}, {escape_sql(name_clean)}, {escape_sql(reg_clean)}, {escape_sql(name_clean)}, {escape_sql(meta['year'])}, {escape_sql(meta['cohort'])}, 'Artificial Intelligence & Data Science', {escape_sql(meta['class'])}, {escape_sql(adv_clean)}, {escape_sql(email)}, NOW(), NOW())"
            )

            profile_rows.append(
                f"({escape_sql(reg_clean)}, {escape_sql(reg_clean)}, {escape_sql(name_clean)}, {escape_sql(email)}, 'Artificial Intelligence & Data Science', {escape_sql(meta['year'])}, {escape_sql(meta['year_sem'])}, {escape_sql(meta['class'])}, {escape_sql(adv_clean)}, NOW(), NOW())"
            )

with open(r'c:\Users\erohi\OneDrive\Desktop\Virtual-Labs\scripts\seed_students.sql', 'w', encoding='utf-8') as f:
    f.write('INSERT INTO public.students (student_id, s_no, essl_code, name, register_number, password, year, cohort, department, class_name, advisor, email, created_at, last_active)\nVALUES\n')
    f.write(',\n'.join(student_rows))
    f.write('\nON CONFLICT (register_number) DO UPDATE SET\n')
    f.write('  name = EXCLUDED.name,\n  password = EXCLUDED.password,\n  year = EXCLUDED.year,\n  cohort = EXCLUDED.cohort,\n  class_name = EXCLUDED.class_name,\n  advisor = EXCLUDED.advisor,\n  last_active = NOW();\n\n')

    f.write('INSERT INTO public.profiles (id, register_number, name, email, department, year, year_semester, class_name, advisor, created_at, last_active)\nVALUES\n')
    f.write(',\n'.join(profile_rows))
    f.write('\nON CONFLICT (register_number) DO UPDATE SET\n')
    f.write('  name = EXCLUDED.name,\n  year = EXCLUDED.year,\n  year_semester = EXCLUDED.year_semester,\n  class_name = EXCLUDED.class_name,\n  advisor = EXCLUDED.advisor,\n  last_active = NOW();\n')

print(f"Generated scripts/seed_students.sql with {len(student_rows)} students and profiles successfully.")
